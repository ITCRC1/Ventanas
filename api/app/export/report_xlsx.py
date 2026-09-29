"""Excel de un reporte: encabezado formal + tabla formateada.

Pensado para mandar afuera (junta, dueños), así que la hoja lleva título, fecha
de generación y los montos con formato contable — no un volcado crudo de la vista.
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
RULE = "C9D2DF"
ZEBRA = "F2F5F9"

_MONEY = "#,##0.00"
_INT = "#,##0"
_DATE = "yyyy-mm-dd"


def _fmt(kind: str) -> str | None:
    return {"money": _MONEY, "int": _INT, "pct": "0.00", "date": _DATE}.get(kind)


def build(
    *,
    title: str,
    blurb: str,
    columns: list[dict[str, str]],
    rows: list[dict[str, Any]],
    cutoff: date | None = None,
) -> bytes:
    """Devuelve el .xlsx del reporte. `columns` es [{key, label, kind}]."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"
    ncol = max(len(columns), 1)

    # --- encabezado ---------------------------------------------------------
    ws.cell(1, 1, "VENTANAS DEVELOPMENT").font = Font(bold=True, size=9, color="44546A")
    ws.cell(2, 1, title).font = Font(bold=True, size=15, color=NAVY)
    ws.cell(3, 1, blurb).font = Font(size=9, color="44546A")
    pie = f"Generated {datetime.now().strftime('%B %d, %Y')}"
    if cutoff:
        pie += f"   ·   Cut-off {cutoff.isoformat()}"
    ws.cell(4, 1, pie).font = Font(size=9, italic=True, color="44546A")
    for r in (1, 2, 3, 4):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)

    # --- cabecera de la tabla ------------------------------------------------
    HEAD = 6
    relleno = PatternFill("solid", fgColor=NAVY)
    borde = Border(bottom=Side("thin", color=RULE))
    for j, c in enumerate(columns, start=1):
        celda = ws.cell(HEAD, j, c["label"])
        celda.font = Font(bold=True, color="FFFFFF", size=9)
        celda.fill = relleno
        celda.alignment = Alignment(
            horizontal="right" if c["kind"] in ("money", "int", "pct") else "left",
            vertical="center", wrap_text=True,
        )
    ws.row_dimensions[HEAD].height = 28

    # --- datos ---------------------------------------------------------------
    cebra = PatternFill("solid", fgColor=ZEBRA)
    for i, fila in enumerate(rows):
        r = HEAD + 1 + i
        for j, c in enumerate(columns, start=1):
            v = fila.get(c["key"])
            if isinstance(v, Decimal):
                v = float(v)
            celda = ws.cell(r, j, v)
            nf = _fmt(c["kind"])
            if nf:
                celda.number_format = nf
            celda.font = Font(size=9)
            celda.border = borde
            if i % 2:
                celda.fill = cebra
            if c["kind"] in ("money", "int", "pct"):
                celda.alignment = Alignment(horizontal="right")

    # --- totales de las columnas de dinero -----------------------------------
    if rows:
        tot = HEAD + 1 + len(rows)
        ws.cell(tot, 1, "TOTAL").font = Font(bold=True, size=9)
        for j, c in enumerate(columns, start=1):
            if c["kind"] != "money":
                continue
            col = get_column_letter(j)
            celda = ws.cell(tot, j, f"=SUM({col}{HEAD + 1}:{col}{tot - 1})")
            celda.number_format = _MONEY
            celda.font = Font(bold=True, size=9)
            celda.alignment = Alignment(horizontal="right")
        for j in range(1, ncol + 1):
            ws.cell(tot, j).border = Border(top=Side("medium", color=NAVY))

    # --- ancho de columna ----------------------------------------------------
    for j, c in enumerate(columns, start=1):
        largo = len(c["label"])
        for fila in rows[:200]:
            largo = max(largo, len(str(fila.get(c["key"]) or "")))
        ws.column_dimensions[get_column_letter(j)].width = min(max(largo + 2, 10), 46)

    ws.freeze_panes = ws.cell(HEAD + 1, 1)
    ws.print_title_rows = f"{HEAD}:{HEAD}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
